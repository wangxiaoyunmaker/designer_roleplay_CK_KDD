#!/usr/bin/env python3
"""Import the final Q1-Q6 cleaning workbook into Route A's frozen CSV contract."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import tempfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


QUESTIONS = {f"Q{index}" for index in range(1, 7)}
SOURCE_MATCH_FIELDS = (
    "question_id",
    "answer_value",
    "duration_seconds",
    "device_type",
    "submitted_at",
)
FIELDS = (
    "record_id",
    "anonymous_user_id",
    "question_id",
    "raw_text_sha256",
    "raw_text",
    "cleaning_status",
    "clean_text",
    "editor_id",
    "review_note",
)


class ImportContractError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        return [
            {str(key).lstrip("\ufeff"): "" if value is None else str(value) for key, value in row.items()}
            for row in csv.DictReader(handle)
        ]


def load_anonymous_mapping(path: Path) -> dict[str, str]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    mapping = payload.get("anonymous_to_raw")
    if not isinstance(mapping, dict) or not mapping:
        raise ImportContractError("mapping must contain nonempty anonymous_to_raw")
    normalized = {str(anonymous): str(raw) for anonymous, raw in mapping.items()}
    if len(set(normalized.values())) != len(normalized):
        raise ImportContractError("anonymous_to_raw contains duplicate raw IDs")
    return normalized


def load_cleaned_workbook(path: Path, sheet_name: str) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportContractError("openpyxl is required to read the cleaning workbook") from exc
    workbook = load_workbook(path, read_only=True, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ImportContractError(f"missing worksheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise ImportContractError("cleaning workbook is empty") from exc
    headers = [canonical_cell(value).strip() for value in header_row]
    required = {"user_id", "text_supplement", *SOURCE_MATCH_FIELDS}
    if not required.issubset(headers):
        raise ImportContractError(f"workbook columns missing: {sorted(required - set(headers))}")
    indexes = {name: headers.index(name) for name in required}
    rows: list[dict[str, str]] = []
    for excel_row, values in enumerate(iterator, start=2):
        question_id = canonical_cell(values[indexes["question_id"]]).strip()
        if question_id not in QUESTIONS:
            continue
        displayed_id = canonical_cell(values[indexes["user_id"]]).strip()
        clean_text = canonical_cell(values[indexes["text_supplement"]])
        if not displayed_id:
            raise ImportContractError(f"blank user_id at Excel row {excel_row}")
        if not clean_text.strip():
            raise ImportContractError(f"blank cleaned text at Excel row {excel_row}")
        if displayed_id.startswith("=") or clean_text.startswith("="):
            raise ImportContractError(f"formula-like value at Excel row {excel_row}")
        rows.append({
            "displayed_user_id": displayed_id,
            "clean_text": clean_text,
            **{
                field: canonical_cell(values[indexes[field]])
                for field in SOURCE_MATCH_FIELDS
            },
        })
    return rows


def _source_fingerprint(row: dict[str, str]) -> tuple[str, ...]:
    values = []
    for field in SOURCE_MATCH_FIELDS:
        value = canonical_cell(row.get(field, ""))
        if field == "duration_seconds":
            try:
                value = format(Decimal(value).normalize(), "f")
            except InvalidOperation:
                pass
        values.append(value)
    return tuple(values)


def reconcile_workbook_rows(
    workbook_rows: Iterable[dict[str, str]],
    source_rows: Iterable[dict[str, str]],
) -> tuple[list[dict[str, str]], dict[str, int]]:
    source: dict[tuple[str, ...], dict[str, str]] = {}
    for row in source_rows:
        if row.get("question_id") not in QUESTIONS:
            continue
        key = _source_fingerprint(row)
        if key in source:
            raise ImportContractError("source task CSV has an ambiguous immutable fingerprint")
        source[key] = row

    reconciled: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    displayed_to_source: dict[str, set[str]] = {}
    mismatched_rows = 0
    for row in workbook_rows:
        key = _source_fingerprint(row)
        if key not in source:
            raise ImportContractError(
                f"workbook row does not match the frozen task source: {row.get('question_id')}"
            )
        source_row = source[key]
        raw_user_id = canonical_cell(source_row.get("user_id", "")).strip()
        displayed = row["displayed_user_id"]
        if not raw_user_id:
            raise ImportContractError("source task CSV contains blank user_id")
        displayed_to_source.setdefault(displayed, set()).add(raw_user_id)
        if displayed != raw_user_id:
            mismatched_rows += 1
        record_key = (raw_user_id, row["question_id"])
        if record_key in seen:
            raise ImportContractError(f"duplicate reconciled workbook record: {record_key}")
        seen.add(record_key)
        reconciled.append({
            "raw_user_id": raw_user_id,
            "question_id": row["question_id"],
            "clean_text": row["clean_text"],
        })

    inconsistent = {
        displayed: sorted(values)
        for displayed, values in displayed_to_source.items()
        if len(values) != 1
    }
    if inconsistent:
        raise ImportContractError(
            f"one workbook user ID maps to multiple source IDs: {list(inconsistent)[:5]}"
        )
    if len(reconciled) != len(source):
        raise ImportContractError(
            f"workbook/source Q1-Q6 coverage differs: workbook={len(reconciled)}, source={len(source)}"
        )
    return reconciled, {
        "identity_mismatch_rows_reconciled": mismatched_rows,
        "identity_mismatch_users_reconciled": sum(
            displayed != next(iter(values))
            for displayed, values in displayed_to_source.items()
        ),
    }


def build_formal_rows(
    template_rows: Iterable[dict[str, str]],
    anonymous_to_raw: dict[str, str],
    cleaned_rows: Iterable[dict[str, str]],
    *,
    editor_id: str,
    workbook_sha256: str,
) -> tuple[list[dict[str, str]], dict[str, int]]:
    if not editor_id.strip():
        raise ImportContractError("editor_id must be nonblank")
    cleaned: dict[tuple[str, str], str] = {}
    for row in cleaned_rows:
        key = (row["raw_user_id"], row["question_id"])
        if key in cleaned:
            raise ImportContractError(f"duplicate workbook record: {key}")
        cleaned[key] = row["clean_text"]

    output: list[dict[str, str]] = []
    expected_keys: set[tuple[str, str]] = set()
    counts: Counter[str] = Counter()
    seen_records: set[str] = set()
    for source in template_rows:
        if set(source) != set(FIELDS):
            raise ImportContractError("template columns do not match the frozen cleaning contract")
        record_id = source["record_id"]
        if record_id in seen_records:
            raise ImportContractError(f"duplicate template record_id: {record_id}")
        seen_records.add(record_id)
        anonymous = source["anonymous_user_id"]
        question_id = source["question_id"]
        if anonymous not in anonymous_to_raw:
            raise ImportContractError(f"template participant missing from private mapping: {anonymous}")
        key = (anonymous_to_raw[anonymous], question_id)
        expected_keys.add(key)
        if key not in cleaned:
            raise ImportContractError(f"cleaned workbook is missing record: {record_id}")
        clean_text = cleaned[key]
        status = "KEEP" if clean_text == source["raw_text"] else "EDITED"
        row = dict(source)
        row.update(
            cleaning_status=status,
            clean_text=clean_text,
            editor_id=editor_id,
            review_note=f"source_workbook_sha256={workbook_sha256}",
        )
        output.append(row)
        counts[status] += 1

    extra = sorted(set(cleaned) - expected_keys)
    if extra:
        raise ImportContractError(f"cleaned workbook contains unexpected Q1-Q6 records: {extra[:5]}")
    if len(output) != 972:
        raise ImportContractError(f"expected 972 formal rows, got {len(output)}")
    return output, dict(sorted(counts.items()))


def write_csv_atomic(path: Path, rows: list[dict[str, str]]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", newline="", dir=path.parent, delete=False
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
        temporary = Path(handle.name)
    os.replace(temporary, path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--source-task-csv", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet", default="task_responses_cleaned")
    parser.add_argument("--editor-id", required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_sha256 = sha256_file(args.workbook)
    cleaned, identity = reconcile_workbook_rows(
        load_cleaned_workbook(args.workbook, args.sheet),
        read_csv_rows(args.source_task_csv),
    )
    rows, counts = build_formal_rows(
        read_csv_rows(args.template),
        load_anonymous_mapping(args.mapping),
        cleaned,
        editor_id=args.editor_id,
        workbook_sha256=workbook_sha256,
    )
    write_csv_atomic(args.output, rows)
    print(json.dumps({
        "status": "PASS",
        "output": str(args.output.resolve()),
        "records": len(rows),
        "cleaning_status_counts": counts,
        "identity_reconciliation": identity,
        "workbook_sha256": workbook_sha256,
        "output_sha256": sha256_file(args.output),
    }, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
