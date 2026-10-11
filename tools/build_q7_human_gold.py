#!/usr/bin/env python3
"""Build unit-level Q7 evaluation gold from adjudicated annotations."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ID_ALIASES = {
    "UNKNOWN_ID_待回填": "-DevolS",
    "0": "-wjh040216",
}
LABEL_TO_SEQUENCE = {
    "K": ["K"],
    "C": ["C"],
    "K→C": ["KtoC"],
    "C→K": ["CtoK"],
    "K→K": ["KtoK"],
    "C→C": ["CtoC"],
    "V": ["V"],
    "C→K→C": ["CtoK", "KtoC"],
    "C→K、K→C": ["CtoK", "KtoC"],
}
BASE_COLUMNS = (
    "designer_id",
    "seg",
    "idx",
    "unit_text",
    "标注员1标签",
    "标注员2标签",
    "仲裁后标签",
)
FINAL_GOLD_COLUMN = "金标准"


def nonempty(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_index(header_row: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, raw_name in enumerate(header_row):
        name = clean(raw_name)
        if not name:
            continue
        if name in columns:
            raise ValueError(f"duplicate column header: {name!r}")
        columns[name] = index
    missing = [name for name in BASE_COLUMNS if name not in columns]
    if missing:
        raise ValueError(f"missing required columns: {missing}")
    return columns


def cell(row: tuple[object, ...], columns: dict[str, int], name: str) -> object:
    index = columns[name]
    return row[index] if index < len(row) else None


def normalize_label(label: object, source_row: int) -> tuple[str, list[str]]:
    normalized = clean(label)
    if normalized not in LABEL_TO_SEQUENCE:
        raise ValueError(f"unsupported final label at row {source_row}: {normalized!r}")
    return normalized, list(LABEL_TO_SEQUENCE[normalized])


def build_payload(input_path: Path) -> dict[str, Any]:
    sheet = load_workbook(input_path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("annotation workbook is empty") from exc
    columns = header_index(header_row)
    uses_final_gold = FINAL_GOLD_COLUMN in columns

    records = []
    exclusions = []
    seen_keys = set()
    source_label_counts: dict[str, int] = {}
    for source_row, row in enumerate(rows, start=2):
        if not any(nonempty(value) for value in row):
            continue

        raw_designer_id = clean(cell(row, columns, "designer_id"))
        designer_id = ID_ALIASES.get(raw_designer_id, raw_designer_id)
        segment = clean(cell(row, columns, "seg"))
        raw_unit_index = cell(row, columns, "idx")
        if not designer_id or not segment or not nonempty(raw_unit_index):
            raise ValueError(f"missing evaluation key field at row {source_row}")
        try:
            unit_index = int(raw_unit_index)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"invalid idx at row {source_row}: {raw_unit_index!r}") from exc

        key = (designer_id, segment, unit_index)
        if key in seen_keys:
            raise ValueError(f"duplicate evaluation key: {key}")
        seen_keys.add(key)

        if uses_final_gold:
            final_label = cell(row, columns, FINAL_GOLD_COLUMN)
            if not nonempty(final_label):
                raise ValueError(f"empty {FINAL_GOLD_COLUMN!r} at row {source_row}")
        else:
            label_one = cell(row, columns, "标注员1标签")
            label_two = cell(row, columns, "标注员2标签")
            adjudicated = cell(row, columns, "仲裁后标签")
            final_label = label_one if clean(label_one) == clean(label_two) else adjudicated
            if not nonempty(final_label):
                exclusions.append({
                    "designer_id": designer_id,
                    "segment": segment,
                    "unit_index": unit_index,
                    "reason": "missing_final_adjudication",
                })
                continue

        source_label, sequence = normalize_label(final_label, source_row)
        source_label_counts[source_label] = source_label_counts.get(source_label, 0) + 1
        unit_text = cell(row, columns, "unit_text")
        if not nonempty(unit_text):
            raise ValueError(f"empty unit text at row {source_row}")
        records.append({
            "designer_id": designer_id,
            "segment": segment,
            "unit_index": unit_index,
            "unit_text": clean(unit_text),
            "gold_state_sequence": sequence,
        })

    final_label_rule = (
        "use the authoritative 金标准 column"
        if uses_final_gold
        else "agreement uses the shared label; disagreement uses the final adjudication column"
    )
    return {
        "schema_version": "q7_human_unit_gold_v2",
        "source": {
            "file_name": input_path.name,
            "sha256": sha256(input_path),
            "final_label_rule": final_label_rule,
            "columns": list(columns),
        },
        "state_symbols": ["K", "C", "KtoC", "CtoK", "KtoK", "CtoC", "V"],
        "compound_label_normalization": {
            "C→K→C": ["CtoK", "KtoC"],
            "C→K、K→C": ["CtoK", "KtoC"],
        },
        "id_aliases": ID_ALIASES,
        "source_label_counts": source_label_counts,
        "records": records,
        "exclusions": exclusions,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    payload = build_payload(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "schema_version": payload["schema_version"],
        "records": len(payload["records"]),
        "exclusions": len(payload["exclusions"]),
        "source_sha256": payload["source"]["sha256"],
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
